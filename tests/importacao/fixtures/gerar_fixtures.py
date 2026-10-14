# Gera as planilhas usadas nos testes do leitor (tests/importacao).
# Rodar:  uv run --with openpyxl==3.1.5 python tests/importacao/fixtures/gerar_fixtures.py
#
# - openpyxl.xlsx: gerada por uma biblioteca independente (textos compartilhados,
#   compactação deflate, datas com formato, número, fórmula, booleano, células vazias).
# - estilo_excel.xlsx: montada à mão no formato que o Excel grava (prefixos de
#   namespace, texto com formatação em partes, fonética, data com formato
#   personalizado, sistema de datas 1904, aba oculta).
# - latin1.csv: CSV com ponto e vírgula na codificação antiga do Windows.
import datetime as dt
import pathlib
import zipfile

import openpyxl

PASTA = pathlib.Path(__file__).parent

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Escolas"
ws.append(["Nome da escola", "Código", "Alunos", "Valor", "Vencimento", "Ativa"])
ws.append(["CEI São José", "UE-01", 120, 1234.56, dt.date(2026, 10, 5), True])
ws.append(["Escola \"Aspas\" & <Símbolos>", None, 80.5, None, dt.datetime(2026, 1, 31, 14, 30), False])
ws.append([])
ws.append(["Depois da linha vazia", "UE-03", "=1+2", None, None, None])
ws["C5"] = "=1+2"
ws["E2"].number_format = "dd/mm/yyyy"
ws["E3"].number_format = "dd/mm/yyyy hh:mm"
ws["J7"] = "longe"
segunda = wb.create_sheet("Outra aba")
segunda.append(["x"])
wb.save(PASTA / "openpyxl.xlsx")

CT = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"><Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/><Default Extension="xml" ContentType="application/xml"/><Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/><Override PartName="/xl/worksheets/sheet1.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/><Override PartName="/xl/worksheets/sheet2.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/><Override PartName="/xl/styles.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.styles+xml"/><Override PartName="/xl/sharedStrings.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml"/></Types>"""
RELS = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="xl/workbook.xml"/></Relationships>"""
WB = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<x:workbook xmlns:x="http://schemas.openxmlformats.org/spreadsheetml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"><x:workbookPr date1904="1"/><x:sheets><x:sheet name="Oculta" sheetId="2" state="hidden" r:id="rId2"/><x:sheet name="Contas &amp; Água" sheetId="1" r:id="rId1"/></x:sheets></x:workbook>"""
WB_RELS = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId3" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/styles" Target="styles.xml"/><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" Target="/xl/worksheets/sheet1.xml"/><Relationship Id="rId2" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" Target="worksheets/sheet2.xml"/><Relationship Id="rId4" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/sharedStrings" Target="sharedStrings.xml"/></Relationships>"""
STYLES = r"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<x:styleSheet xmlns:x="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><x:numFmts count="2"><x:numFmt numFmtId="164" formatCode="[$-416]d/mmm/yyyy;@"/><x:numFmt numFmtId="165" formatCode="&quot;R$&quot;\ #,##0.00"/></x:numFmts><x:cellXfs count="4"><x:xf numFmtId="0" fontId="0" fillId="0" borderId="0" xfId="0"/><x:xf numFmtId="164" fontId="0" fillId="0" borderId="0" xfId="0" applyNumberFormat="1"/><x:xf numFmtId="165" fontId="0" fillId="0" borderId="0" xfId="0" applyNumberFormat="1"/><x:xf numFmtId="14" fontId="0" fillId="0" borderId="0" xfId="0" applyNumberFormat="1"/></x:cellXfs></x:styleSheet>"""
SST = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<x:sst xmlns:x="http://schemas.openxmlformats.org/spreadsheetml/2006/main" count="4" uniqueCount="4"><x:si><x:t>Escola</x:t></x:si><x:si><x:r><x:rPr><x:b/></x:rPr><x:t xml:space="preserve">CEI </x:t></x:r><x:r><x:t>Girassol</x:t></x:r><x:rPh sb="0" eb="1"><x:t>FONETICA</x:t></x:rPh></x:si><x:si><x:t>Vencimento</x:t></x:si><x:si><x:t>Valor</x:t></x:si></x:sst>"""
SHEET1 = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<x:worksheet xmlns:x="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><x:sheetData><x:row r="1"><x:c r="A1" t="s"><x:v>0</x:v></x:c><x:c r="B1" t="s"><x:v>2</x:v></x:c><x:c r="C1" t="s"><x:v>3</x:v></x:c></x:row><x:row r="2"><x:c r="A2" t="s"><x:v>1</x:v></x:c><x:c r="B2" s="1"><x:v>44474</x:v></x:c><x:c r="C2" s="2"><x:v>1500.5</x:v></x:c></x:row><x:row r="3"><x:c r="A3" t="inlineStr"><x:is><x:t>Inline &#231;&#xE3;o</x:t></x:is></x:c><x:c r="B3" s="3"><x:v>0</x:v></x:c><x:c r="C3" t="e"><x:v>#N/A</x:v></x:c></x:row></x:sheetData></x:worksheet>"""
SHEET2 = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><sheetData><row r="1"><c r="A1" t="inlineStr"><is><t>escondida</t></is></c></row></sheetData></worksheet>"""

with zipfile.ZipFile(PASTA / "estilo_excel.xlsx", "w", zipfile.ZIP_DEFLATED) as z:
    z.writestr("[Content_Types].xml", CT)
    z.writestr("_rels/.rels", RELS)
    z.writestr("xl/workbook.xml", WB)
    z.writestr("xl/_rels/workbook.xml.rels", WB_RELS)
    z.writestr("xl/styles.xml", STYLES)
    z.writestr("xl/sharedStrings.xml", SST)
    z.writestr("xl/worksheets/sheet1.xml", SHEET1)
    z.writestr("xl/worksheets/sheet2.xml", SHEET2)

(PASTA / "latin1.csv").write_bytes(
    'Escola;Valor;Observações\r\nCEI Ação;"1.234,56";"tem ""aspas""; e ponto e vírgula"\r\nCEI B;10;\r\n'.encode("cp1252")
)
print("ok")

# Planilha "do mundo real" para testar a tela de importação de contas no
# emulador/demonstração: título antes dos cabeçalhos, nomes de coluna
# diferentes do modelo, e erros de propósito.
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Outubro"
ws.append(["Contas do escritório — outubro"])
ws.append([])
ws.append(["Unidade", "Conta", "Mês de referência", "Valor total", "Data de vencimento", "Nº NF", "Obs"])
linhas = [
    ["CEI Pequenos Passos", "Energia", "10/2026", 1234.56, dt.date(2026, 11, 10), "E-100", ""],
    ["UE-02", "Água", "out/2026", "R$ 512,30", "12/11/2026", "A-200", ""],
    ["Arco-Íris", "Dedetização", "Outubro/2026", "890", "20/11/2026", "NF-31", "semestral"],
    ["CEI Inexistente", "CPFL", "10/2026", 100, "10/11/2026", "", ""],
    ["CEI Semear", "CPFL", "10/2026", 0, "10/11/2026", "", ""],
    ["CEI Semear", "SABESP", "10/2026", 70, "31/11/2026", "", ""],
    ["CEI Pequenos Passos", "luz", "outubro/2026", "1.234,56", "10/11/2026", "e-100", "repetida"],
    ["CEI Girassol", "Taxa de limpeza", "10/2026", 300, dt.date(2026, 11, 5), "TL-9", ""],
]
for l in linhas:
    ws.append(l)
for celula in ("E4", "E11"):
    ws[celula].number_format = "dd/mm/yyyy"
wb.save(PASTA / "contas-com-erros.xlsx")
print("ok 2")
